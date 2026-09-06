#!/usr/bin/env python3
"""Fill Ningbo angel-fund attachment 3/4 for the founder. Do not invent teammate names."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

SRC3 = Path("/Users/damowang/Downloads/附件3.种子直投项目经营（商业）计划书（科技团队）-修改.docx")
SRC4 = Path("/Users/damowang/Downloads/附件4.承诺书（科技团队）.docx")
DESK = Path.home() / "Desktop" / "宁波科技大脑-申报附件"
KEEP = Path.home() / "Desktop" / "灵祺种子轮投融资资料-20260906" / "自己留着-投递清单" / "宁波科技大脑-申报附件"

BODY = {
    "行业背景及现状、行业特点、行业规模；发展前景及增长趋势分析与预测。": (
        "本地生活门店已被抖音、美团等平台推着做短视频和达人种草，但经营仍大量停在微信群和表格："
        "组品、找达人、交付、结算拆成几截。公开数据转述：本地商家服务业 2024 年规模超约 13.7 万亿，"
        "商家超约 1,300 万；餐饮线上运营营销支出约 2,420 亿，其中解决方案市场约 691 亿；前五大合计约 3%，"
        "市场极度分散。杭甬温有达人种草或短视频预算的门店，量级约 0.8–1.5 万家。"
        "大模型让「AI 出方案、系统去执行」第一次能用在真实经营里。窗口约 12–24 个月，先把杭甬温做深，不是全国铺开。"
    ),
    "介绍核心团队的教育经历、工作经历、获得奖项/主要成就、重大项目参与情况等。": (
        "项目负责人：曹鑫淼。负责战略、产品与融资，统筹灵祺 AI ERP 与灵祺星选达人履约。"
        "联系电话 15757468650，邮箱 modian@mofangdianai.com。"
        "已推动产品在生产环境上线，取得软件著作权 2 项：灵祺 AI 智能 ERP 商家管理系统 V1.0、"
        "灵祺达人招募小程序软件 V1.0。\n"
        "核心成员 2（请打印前手写补全姓名、学历、职责，勿空签）：________，负责技术交付（ERP、星选与线上系统）。\n"
        "核心成员 3（请打印前手写补全姓名、学历、职责，勿空签）：________，负责运营或商务（达人供给 / 服务商与门店拓展）。\n"
        "编制规划：研发 3–4 人（含现有）、产品运营 1–2 人、商务 2 人、创始团队 2 人。除曹鑫淼外，其余姓名未写入系统材料，须本人核对后手填并签字。"
    ),
    "主营产品（服务）介绍：包括但不限于产品（服务）简介，市场需求分析，市场定位及合理性分析，与国内外同行业其它公司同类产品（服务）的比较等。": (
        "主营产品为本地生活商家 AI ERP「灵祺」，以及达人履约小程序「灵祺星选」。"
        "商家在系统里用 AI 组品、分析店铺、制作短视频，确认后再写入系统；达人报名探店或云剪并回传，两边在同一套系统里。"
        "线上入口：商家 ERP https://cs.mofangdianai.com ，服务商 ERP https://fws.mofangdianai.com ，"
        "星选 https://dr.mofangdianai.com ，微信搜「灵祺星选」。"
        "目标客户：杭甬温餐饮、零售、休闲娱乐等本地生活门店，以及代运营服务商。"
        "与通用聊天助手的差别是商户确认后再写入系统；与只发招募的工具的差别是和经营在同一后台；"
        "与高客单代运营的差别是低价自助、服务商可批量带店；与有赞的差别是自己有达人履约；"
        "与来客、美团后台的差别是多平台一起用，不替代平台交易和资金。"
    ),
    "主要介绍本产品（服务）所采用的核心技术（模式）及其创新性、先进性，及与国内外同行业其它公司同类产品（服务）核心技术的比较等。": (
        "核心不是再做一个聊天窗口，而是把大模型用在经营流程里：AI 生成组品、方案、店铺分析和短视频，"
        "商户确认后再写入 ERP，再由星选完成达人履约和回传。线上接口 212 条，登录、大厅、支付已上线，"
        "服务器连续运行 100 天。截至 2026 年 9 月 6 日，AI 累计调用 6,397 次，7 月高峰 4,286 次。"
        "已取得上述 2 项软著，算法备案材料推进中。与同类产品比，护城河是「ERP 发单 → 星选履约 → 回传 → 续费」打通，拆开就会被收银系统和只发招募的工具替换。"
    ),
    "项目商业模式及盈利模式、目标客户群体、销售模式、成本与毛利分析、供应链情况等。": (
        "收入来源：商家订阅（约 168 元/月，高阶上浮）、履约撮合（约 10–50 元/单）、"
        "达人与 PR 会员（约 9.9 元/月）、服务商席位（1.98–4.98 万元/年）、按 AI 用量和云剪加价。"
        "目标客户是杭甬温本地生活门店和服务商。销售以服务商带客为主，商家直发刚开始。"
        "毛利率目标：订阅 75–85%，履约 35–50%，是目标不是承诺。"
        "不替代抖音、美团的交易和资金，不自建仓储物流。本轮向宁波市天使投资引导基金申请种子直投 200 万元；"
        "整轮种子目标 400 万元，首次交割 300 万即可启动。"
    ),
    "目前市场推广、渠道销售、已有客户的合作订单（包含意向合作）等进展。": (
        "截至 2026 年 9 月 6 日生产环境可复核：ERP 客户 32 家（具名约 14 家，服务商 5 家，含塔塔、极道、杭州意米、意米文化等）；"
        "星选会员 737 人（温州 264、浙江 424），达人库 527 人（中位粉丝 2.2 万，10 万粉以上 170 人）；"
        "招募单 293 条，报名 1,979 次；3 家深度店已同步抖音来客订单 59,643 单、平台成交 733 万（商户成交，不是本公司收入），覆盖 107 个门店点位。"
        "支付通道已接微信、支付宝、抖音支付，确认到账约 0.29 万元，规模化收费尚未起来。"
        "733 万不得写成公司订单额。"
    ),
    "下一步的研发、生产和销售计划，以及未来三年的营业收入、经营成本、净利润预测。": (
        "下一步：把现有用量做成经常性收入。18 个月内，M6 付费门店 150–300 家，M12 付费 800–1,500 家；"
        "资金用于服务商转付费、门店激活、履约交付，稳住现有系统，不重做产品。"
        "未来三年目标预测（按上一阶段滚动调整，不是承诺）：\n"
        "第一年：营业收入 60 万元，经营成本 300 万元，净利润 -240 万元；\n"
        "第二年：营业收入 280 万元，经营成本 350 万元，净利润 -70 万元；\n"
        "第三年：营业收入 800 万元，经营成本 560 万元，净利润 240 万元。\n"
        "第一年仍以验证付费为主，确认收款从约 0.29 万元做起，不按全国 2 万店测算。"
    ),
    "项目潜在技术风险、市场风险、管理风险及应对措施": (
        "付费样本薄：确认到账约 0.29 万元。对策：90 天做厚 3 个能对账的案例，并推动服务商转付费合同。\n"
        "大厅以 PR 代发为主（290/293）。对策：把商家直发做成默认路径。\n"
        "8 月后达人新增放缓。对策：融资用于做深单城供给，不全国投放。\n"
        "平台后台够用化。对策：ERP 和星选绑在一起，多平台一起用，不碰平台资金。\n"
        "平台成交被误读为收入。对策：对外口径写清 733 万是商户成交，不是公司收入。"
    ),
}


HOWTO = """项目申请书（必传）怎么处理

系统红字已经写明：请在「申报管理」里下载、打印申报书，按要求签字盖章，再扫描上传。

这份申请书不是桌面上的 Word，是你在科技大脑里填完表单后，系统生成的申报书。

步骤：
1. 回到「申报管理」，找到本项目，下载「项目申请书」PDF 或 Word。
2. 核一遍：项目名称、申请金额 200 万元、投前估值 2000–2800 万元、公司宁波墨典、负责人曹鑫淼。
3. 打印，项目负责人和核心成员按要求签字，企业已设立的加盖宁波墨典公章。
4. 扫描成 PDF，在附件页点「上传」。

桌面这两份可以先传：
  附件3-商业计划书-已填写.docx
  附件4-承诺书-已填写.docx
承诺书签字前，把另外 2 名核心成员姓名手写进空格，三人都要签字。
承诺书第四条是模板原文（6 个月内在宁波设立企业）。公司已是宁波墨典，按原文签字即可，不要改模板法律条款。

不要上传：资金评估、机构投递清单、各机构邮件模板。
"""


def set_run_text(run, text: str) -> None:
    run.text = text
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")


def fill_commitment(src: Path, dest: Path) -> None:
    doc = Document(str(src))
    p = doc.paragraphs[2]
    p.runs[1].text = "[曹鑫淼]"
    p.runs[3].text = "[曹鑫淼、________、________]"
    date = doc.paragraphs[11]
    date.runs[0].text = "2026 年"
    date.runs[1].text = " "
    date.runs[2].text = "9 月"
    date.runs[3].text = " "
    date.runs[4].text = "6 日"
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))


def fill_plan(src: Path, dest: Path) -> None:
    doc = Document(str(src))
    for para in doc.paragraphs:
        key = para.text.strip()
        if key not in BODY:
            continue
        if not para.runs:
            para.add_run(BODY[key])
            continue
        first = para.runs[0]
        first.text = BODY[key]
        first.font.size = Pt(16)
        first.font.name = "Times New Roman"
        rpr = first._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "宋体")
        for extra in para.runs[1:]:
            extra.text = ""
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))


def main() -> None:
    if not SRC3.exists() or not SRC4.exists():
        raise SystemExit("missing source docx in Downloads")
    DESK.mkdir(parents=True, exist_ok=True)
    KEEP.mkdir(parents=True, exist_ok=True)
    plan = DESK / "附件3-商业计划书-已填写.docx"
    letter = DESK / "附件4-承诺书-已填写.docx"
    fill_plan(SRC3, plan)
    fill_commitment(SRC4, letter)
    (DESK / "00-项目申请书怎么传.txt").write_text(HOWTO, encoding="utf-8")
    for src in (plan, letter, DESK / "00-项目申请书怎么传.txt"):
        (KEEP / src.name).write_bytes(src.read_bytes())
    print("DESK", DESK)
    print("KEEP", KEEP)


if __name__ == "__main__":
    main()
